"""The extraction queue's storage: a human sheet, a machine ledger, a heartbeat.

The sheet is what a person edits -- one row per chapter, grouped by standard and
subject, carrying the PDF link. It is an ordinary .xlsx (or .csv), so it can be
filled in Excel with no Google account and no API key.

The ledger is what the runner trusts. It exists because the sheet cannot be
relied on at 3 a.m.: on Windows an open Excel window holds an exclusive lock and
every write to the workbook fails with PermissionError. If the sheet were the
only record, one forgotten open window would lose a whole night of bookkeeping,
and the next run would redo every chapter. So:

    ledger.jsonl   append-only, one JSON object per state change, never locked.
                   This is the resume record.
    the workbook   best-effort mirror, refreshed after every chapter so the
                   morning view is readable. If it is locked the runner logs
                   once and carries on; --sync-sheet replays the ledger into it.

Both are keyed the same way: (board, standard, subject, document_type, chapter).
That tuple is also how document_extractions rows are matched, so a row in the
sheet, a line in the ledger and a row in the database all name the same chapter.

A third file, run_state.json, is the heartbeat: what the run is doing right now,
rewritten on every state change. It exists so the web UI can show a live night
without talking to the runner at all -- the runner is a detached process that
outlives the API, so shared memory was never an option. The same directory holds
the STOP flag, which is how the morning stop reaches a process nobody has a
handle on.
"""

from __future__ import annotations

import contextlib
import csv
import json
import logging
import os
import tempfile
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator

logger = logging.getLogger("queue.sheet")

QUEUE_SHEET_NAME = "queue"

# Columns a person fills in. This order is the column order of a generated book.
INPUT_COLUMNS = [
    "enabled",
    "board",
    "standard",
    "subject_name",
    "chapter_number",
    "document_title",
    "document_type",
    "syear",
    "sub_institute_id",
    "pdf_url",
]

# Columns the runner writes back. Do not edit these by hand -- they are
# overwritten from the ledger after every chapter.
OUTPUT_COLUMNS = [
    "status",
    "extraction_id",
    "pages",
    "md_chars",
    "images",
    "seconds",
    "attempts",
    "last_error",
    "finished_at",
]

COLUMNS = INPUT_COLUMNS + OUTPUT_COLUMNS

# A row in one of these states is finished and is not picked up again without
# --force.
TERMINAL = frozenset({"done", "skipped"})

# 'extracted' means MinerU succeeded but the database write did not. The
# markdown is already in the on-disk extraction cache, so a re-run re-persists
# it in seconds without spending the CPU again. Deliberately NOT terminal.
RETRYABLE = frozenset({"", "pending", "failed", "running", "extracted"})


def _norm_header(value: Any) -> str:
    """Forgive the ways a person types a column name."""
    return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _int_or_none(value: Any) -> int | None:
    text = _clean(value)
    if not text:
        return None
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return None


def _float_or_none(value: Any) -> float | None:
    text = _clean(value)
    if not text:
        return None
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


@dataclass
class QueueRow:
    """One chapter to extract."""

    # Where this row sits in the workbook: 1-based and counting the header, so
    # it can be written straight back to the same cells.
    sheet_row: int = 0

    enabled: str = "yes"
    board: str = "CBSE"
    standard: int | None = None
    subject_name: str = ""
    chapter_number: int | None = None
    document_title: str = ""
    document_type: str = "Chapter"
    syear: int | None = None
    sub_institute_id: int | None = None
    pdf_url: str = ""

    status: str = "pending"
    extraction_id: int | None = None
    pages: int | None = None
    md_chars: int | None = None
    images: int | None = None
    seconds: float | None = None
    attempts: int = 0
    last_error: str = ""
    finished_at: str = ""

    @property
    def key(self) -> str:
        """Stable identity, independent of where the row sits in the sheet."""
        return "|".join(
            [
                _clean(self.board).upper(),
                _clean(self.standard),
                _clean(self.subject_name).lower(),
                _clean(self.document_type).lower(),
                _clean(self.chapter_number),
            ]
        )

    @property
    def subject_key(self) -> tuple[str, str, str]:
        """The grouping the runner walks: one subject finishes before the next."""
        return (
            _clean(self.board).upper(),
            _clean(self.standard),
            _clean(self.subject_name).lower(),
        )

    @property
    def label(self) -> str:
        return (
            f"{self.board} std{_clean(self.standard) or '?'} "
            f"{self.subject_name or '?'} ch{_clean(self.chapter_number) or '?'}"
        )

    def is_enabled(self) -> bool:
        return _clean(self.enabled).lower() not in {"no", "n", "false", "0", "off", "skip"}

    def is_runnable(self) -> bool:
        """Should a normal (non-force) run pick this row up?"""
        return self.is_enabled() and _clean(self.status).lower() in RETRYABLE

    def validate(self) -> str:
        """Empty string if this row can run; otherwise why it cannot."""
        if not _clean(self.pdf_url):
            return "pdf_url is empty"
        if self.standard is None:
            return "standard is empty"
        if not _clean(self.subject_name):
            return "subject_name is empty"
        if self.chapter_number is None and _clean(self.document_type).lower() == "chapter":
            return "chapter_number is empty"
        if not _clean(self.document_title):
            return "document_title is empty"
        return ""


def _row_from_mapping(data: dict[str, Any], sheet_row: int) -> QueueRow:
    return QueueRow(
        sheet_row=sheet_row,
        enabled=_clean(data.get("enabled")) or "yes",
        board=_clean(data.get("board")) or "CBSE",
        standard=_int_or_none(data.get("standard")),
        subject_name=_clean(data.get("subject_name")) or _clean(data.get("subject")),
        chapter_number=_int_or_none(data.get("chapter_number")),
        document_title=_clean(data.get("document_title")) or _clean(data.get("chapter_name")),
        document_type=_clean(data.get("document_type")) or "Chapter",
        syear=_int_or_none(data.get("syear")),
        sub_institute_id=_int_or_none(data.get("sub_institute_id")),
        pdf_url=_clean(data.get("pdf_url")),
        status=_clean(data.get("status")).lower() or "pending",
        extraction_id=_int_or_none(data.get("extraction_id")),
        pages=_int_or_none(data.get("pages")),
        md_chars=_int_or_none(data.get("md_chars")),
        images=_int_or_none(data.get("images")),
        seconds=_float_or_none(data.get("seconds")),
        attempts=_int_or_none(data.get("attempts")) or 0,
        last_error=_clean(data.get("last_error")),
        finished_at=_clean(data.get("finished_at")),
    )


# --------------------------------------------------------------------------
# reading
# --------------------------------------------------------------------------

def load(path: Path) -> list[QueueRow]:
    """Read every row of the queue sheet, in sheet order."""
    if not path.exists():
        raise FileNotFoundError(
            f"Queue sheet not found: {path}\n"
            "Create one with: python -m scripts.build_extraction_sheet --standard 10"
        )
    if path.suffix.lower() in {".csv", ".txt"}:
        return _load_csv(path)
    return _load_xlsx(path)


def _load_csv(path: Path) -> list[QueueRow]:
    rows: list[QueueRow] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for index, raw in enumerate(csv.DictReader(handle), start=2):
            data = {_norm_header(k): v for k, v in raw.items()}
            if not any(_clean(v) for v in data.values()):
                continue
            rows.append(_row_from_mapping(data, index))
    return rows


def _load_xlsx(path: Path) -> list[QueueRow]:
    from openpyxl import load_workbook

    # data_only returns a formula's cached value rather than its text, so a
    # pdf_url built with CONCAT still reads as a URL.
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = (
            workbook[QUEUE_SHEET_NAME]
            if QUEUE_SHEET_NAME in workbook.sheetnames
            else workbook.active
        )
        iterator: Iterator[tuple[Any, ...]] = sheet.iter_rows(values_only=True)
        try:
            header = [_norm_header(cell) for cell in next(iterator)]
        except StopIteration:
            return []

        rows: list[QueueRow] = []
        for index, values in enumerate(iterator, start=2):
            if not any(_clean(v) for v in values):
                continue
            rows.append(_row_from_mapping(dict(zip(header, values)), index))
        return rows
    finally:
        workbook.close()


# --------------------------------------------------------------------------
# writing back
# --------------------------------------------------------------------------

class SheetWriter:
    """Mirrors run state into the workbook, tolerating a locked file.

    Excel takes an exclusive lock on an open .xlsx. A runner that treated that
    as fatal would die at the first chapter of the night because someone left
    the sheet open, so a failed write is logged once and then suppressed and the
    ledger carries on as the real record.
    """

    def __init__(self, path: Path) -> None:
        self.path = path
        self.enabled = path.suffix.lower() not in {".csv", ".txt"}
        self._warned = False
        self.writes = 0
        self.failures = 0

    def update(self, rows: Iterable[QueueRow]) -> bool:
        """Write the output columns of the given rows. True if the file changed."""
        rows = [r for r in rows if r.sheet_row > 0]
        if not rows or not self.enabled:
            return False
        try:
            self._write(rows)
            self.writes += 1
            self._warned = False
            return True
        except PermissionError:
            self.failures += 1
            if not self._warned:
                logger.warning(
                    "Queue sheet is locked (open in Excel?) -- progress is still "
                    "recorded in the ledger. Close it, or run --sync-sheet later."
                )
                self._warned = True
            return False
        except Exception:
            self.failures += 1
            if not self._warned:
                logger.warning("Could not write the queue sheet", exc_info=True)
                self._warned = True
            return False

    def _write(self, rows: list[QueueRow]) -> None:
        from openpyxl import load_workbook

        workbook = load_workbook(self.path)
        try:
            sheet = (
                workbook[QUEUE_SHEET_NAME]
                if QUEUE_SHEET_NAME in workbook.sheetnames
                else workbook.active
            )
            header = {
                _norm_header(cell.value): cell.column
                for cell in next(sheet.iter_rows(min_row=1, max_row=1))
            }
            for row in rows:
                for column in OUTPUT_COLUMNS:
                    position = header.get(column)
                    if position is None:
                        continue
                    value = getattr(row, column)
                    if column == "seconds" and value is not None:
                        value = round(float(value), 1)
                    if column == "last_error" and value:
                        # A MinerU traceback runs to thousands of characters and
                        # Excel caps a cell at 32,767. Keep the useful head.
                        value = str(value)[:500]
                    sheet.cell(row=row.sheet_row, column=position, value=value)
            self._save_atomically(workbook)
        finally:
            workbook.close()

    def _save_atomically(self, workbook: Any) -> None:
        """Never leave a half-written workbook behind.

        openpyxl rewrites the whole file; a crash partway through -- and this
        runs unattended overnight -- would truncate the queue itself. Writing to
        a sibling temp file and replacing is atomic on NTFS.
        """
        handle, temp_name = tempfile.mkstemp(
            dir=str(self.path.parent), prefix=".queue-", suffix=".xlsx"
        )
        os.close(handle)
        temp_path = Path(temp_name)
        try:
            workbook.save(temp_path)
            os.replace(temp_path, self.path)
        except Exception:
            temp_path.unlink(missing_ok=True)
            raise


# --------------------------------------------------------------------------
# the ledger
# --------------------------------------------------------------------------

@dataclass
class LedgerEntry:
    key: str
    status: str
    label: str = ""
    extraction_id: int | None = None
    pages: int | None = None
    md_chars: int | None = None
    images: int | None = None
    seconds: float | None = None
    attempts: int = 0
    error: str = ""
    at: str = field(
        default_factory=lambda: datetime.now(timezone.utc).isoformat(timespec="seconds")
    )
    run_id: str = ""


class Ledger:
    """Append-only run record. The file the runner actually trusts."""

    def __init__(self, path: Path, run_id: str = "") -> None:
        self.path = path
        self.run_id = run_id
        self.path.parent.mkdir(parents=True, exist_ok=True)

    def append(self, entry: LedgerEntry) -> None:
        entry.run_id = entry.run_id or self.run_id
        try:
            with self.path.open("a", encoding="utf-8") as handle:
                handle.write(json.dumps(asdict(entry), ensure_ascii=False) + "\n")
                # Flush and fsync every line. A run killed by OOM, a power cut
                # or Ctrl-C must resume from what it had already reported, not
                # from whatever happened to survive in a buffer.
                handle.flush()
                os.fsync(handle.fileno())
        except Exception:
            logger.warning("Could not append to the ledger", exc_info=True)

    def latest(self) -> dict[str, LedgerEntry]:
        """The most recent entry per chapter key."""
        state: dict[str, LedgerEntry] = {}
        if not self.path.exists():
            return state
        fields = set(LedgerEntry.__annotations__)
        with self.path.open("r", encoding="utf-8") as handle:
            for line in handle:
                line = line.strip()
                if not line:
                    continue
                try:
                    data = json.loads(line)
                except json.JSONDecodeError:
                    # A line torn in half by a hard kill. Skip it; the previous
                    # entry for that chapter is still good.
                    continue
                key = data.get("key")
                if not key:
                    continue
                state[key] = LedgerEntry(
                    **{k: v for k, v in data.items() if k in fields}
                )
        return state


def apply_ledger(rows: list[QueueRow], state: dict[str, LedgerEntry]) -> int:
    """Carry ledger state onto sheet rows. Returns how many rows changed.

    Used at startup -- so a sheet whose write-back was blocked by an open Excel
    window still resumes correctly -- and by --sync-sheet.
    """
    changed = 0
    for row in rows:
        entry = state.get(row.key)
        if entry is None:
            continue
        # 'running' in the ledger means the process died mid-chapter. Report it
        # as failed rather than claiming work is still in flight.
        status = "failed" if entry.status == "running" else entry.status
        if (
            row.status == status
            and row.extraction_id == entry.extraction_id
            and row.attempts == entry.attempts
        ):
            continue
        row.status = status
        row.extraction_id = entry.extraction_id
        row.pages = entry.pages
        row.md_chars = entry.md_chars
        row.images = entry.images
        row.seconds = entry.seconds
        row.attempts = entry.attempts
        row.last_error = entry.error
        row.finished_at = entry.at if status not in {"pending", "running"} else ""
        changed += 1
    return changed


# --------------------------------------------------------------------------
# the heartbeat and the stop flag
# --------------------------------------------------------------------------

class RunState:
    """What the run is doing right now, on disk where anyone can read it.

    The runner is launched detached so a night survives the API restarting, an
    IDE closing, or a terminal being shut. That rules out asking it anything
    directly, so it publishes instead: a small JSON file rewritten whenever
    something changes, and a STOP flag it checks between chapters.

    Written the same atomic way as the workbook. A reader that catches the file
    mid-write would show a half-parsed night, and this file is the only thing
    the morning status page has to go on.
    """

    def __init__(self, path: Path) -> None:
        self.path = path
        # The stop flag is named after the state file, so two sheets running
        # side by side each have their own: stopping Class 9 must not also stop
        # Class 10. run_state.json keeps the original bare STOP.
        stem = path.stem.removesuffix(".state")
        self.stop_path = (
            path.parent / "STOP" if path.name == "run_state.json" else path.parent / f"{stem}.STOP"
        )
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.data: dict[str, Any] = {}

    # -- writing (the runner) --------------------------------------------

    def publish(self, **fields: Any) -> None:
        self.data.update(fields)
        self.data["updated_at"] = datetime.now().isoformat(timespec="seconds")
        temp_name = ""
        try:
            handle, temp_name = tempfile.mkstemp(
                dir=str(self.path.parent), prefix=".state-", suffix=".json"
            )
            with os.fdopen(handle, "w", encoding="utf-8") as stream:
                json.dump(self.data, stream, ensure_ascii=False, indent=2, default=str)
            os.replace(temp_name, self.path)
        except Exception:
            # The heartbeat is a convenience for the UI. It must never be able
            # to take down the extraction it is reporting on -- but it must not
            # litter the queue directory either, and this runs on every chapter.
            if temp_name:
                with contextlib.suppress(OSError):
                    Path(temp_name).unlink()
            logger.debug("Could not publish run state", exc_info=True)

    # -- reading (the API) ------------------------------------------------

    def read(self) -> dict[str, Any]:
        if not self.path.exists():
            return {}
        try:
            return json.loads(self.path.read_text(encoding="utf-8"))
        except Exception:
            return {}

    # -- the stop flag ----------------------------------------------------

    def request_stop(self, requested_by: str = "user") -> None:
        """Ask the run to finish its current chapters and stop.

        A file rather than a signal: the API process did not fork the runner and
        on Windows has no clean way to interrupt it, and a half-killed MinerU
        leaves a stub row claiming to be extracting forever.
        """
        self.stop_path.parent.mkdir(parents=True, exist_ok=True)
        self.stop_path.write_text(
            json.dumps(
                {
                    "requested_by": requested_by,
                    "at": datetime.now().isoformat(timespec="seconds"),
                }
            ),
            encoding="utf-8",
        )

    def stop_requested(self) -> bool:
        return self.stop_path.exists()

    def clear_stop(self) -> None:
        self.stop_path.unlink(missing_ok=True)


def companion_paths(sheet: Path) -> tuple[Path, Path]:
    """The ledger and heartbeat that belong to one sheet.

    Two sheets -- one per class, run on two machines -- must not share a ledger.
    If they did, each would read the other's rows as already handled and the
    heartbeats would overwrite each other, so the page would show whichever run
    published last. Naming them after the sheet makes collision impossible and
    keeps a sheet's whole history with it if the file is copied elsewhere.

    The original single-sheet layout is preserved: extraction_queue.xlsx keeps
    ledger.jsonl and run_state.json, so runs recorded before this existed are
    still found.
    """
    if sheet.stem == "extraction_queue":
        return sheet.parent / "ledger.jsonl", sheet.parent / "run_state.json"
    return (
        sheet.parent / f"{sheet.stem}.ledger.jsonl",
        sheet.parent / f"{sheet.stem}.state.json",
    )
