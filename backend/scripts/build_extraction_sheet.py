"""Create or extend the local extraction queue workbook, standard by subject.

The extraction form asks a person for board, standard, subject, chapter number,
chapter title and a PDF link, then extracts one chapter. Forty chapters means
forty rounds of that. This builds the same information as a spreadsheet instead,
and `run_extraction_queue.py` reads it unattended.

The chapter list is not typed by hand. It comes from `chapter_master`, which is
the same table `_map_ids` later matches an extraction against -- so the titles
and chapter numbers in the sheet are, by construction, the ones that will
resolve to a chapter_id rather than quietly landing as NULL.

    # what spellings exist for a class (read-only, no writes)
    python -m scripts.build_extraction_sheet --list --standard 10

    # one subject's chapters, titles filled in, pdf_url left for you
    python -m scripts.build_extraction_sheet --standard 10 --subject Science

    # every subject the class has
    python -m scripts.build_extraction_sheet --standard 10 --all-subjects

    # a class whose chapters are not in chapter_master yet: 15 blank rows
    python -m scripts.build_extraction_sheet --standard 9 --subject Maths --blank 15

Re-running is safe. Existing rows are never touched -- only chapters not already
in the sheet are appended -- so a night's results survive a rebuild.
"""

from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from sqlalchemy import text  # noqa: E402

from app.db.mariadb import SessionLocal, init_mariadb  # noqa: E402
from app.utils.config import settings  # noqa: E402
from scripts import queue_sheet as qs  # noqa: E402

logging.basicConfig(level=logging.INFO, format="%(levelname)-7s %(message)s")
logger = logging.getLogger("sheet")

DEFAULT_SHEET = Path(__file__).resolve().parents[1] / "queue" / "extraction_queue.xlsx"

HEADER_FILL = "FF1F3864"
BAND_FILL = "FFF2F6FC"
COLUMN_WIDTHS = {
    "enabled": 9,
    "board": 12,
    "standard": 9,
    "subject_name": 22,
    "chapter_number": 9,
    "document_title": 46,
    "document_type": 15,
    "syear": 8,
    "sub_institute_id": 9,
    "pdf_url": 62,
    "status": 11,
    "extraction_id": 12,
    "pages": 7,
    "md_chars": 10,
    "images": 7,
    "seconds": 9,
    "attempts": 9,
    "last_error": 44,
    "finished_at": 21,
}


# --------------------------------------------------------------------------
# reading the masters
# --------------------------------------------------------------------------

def _tenant(board: str, override: int | None) -> int:
    return override if override is not None else settings.tenant_for_board(board)


def standards(tenant: int) -> dict[str, int]:
    """name -> id, for the classes this tenant has."""
    with SessionLocal() as db:
        rows = db.execute(
            text("SELECT id, name FROM standard WHERE sub_institute_id = :t ORDER BY id"),
            {"t": tenant},
        ).fetchall()
    return {str(r[1]).strip(): int(r[0]) for r in rows}


def subjects_for(tenant: int, standard_id: int) -> dict[str, int]:
    """subject_name -> id, for subjects that actually have chapters in this class.

    Joined through chapter_master rather than read from `subject` directly: the
    subject table is tenant-wide, so a plain read offers Class 10 the subjects
    of every other class too.
    """
    with SessionLocal() as db:
        rows = db.execute(
            text(
                """
                SELECT s.id, s.subject_name, COUNT(c.id) AS chapters
                  FROM subject s
                  JOIN chapter_master c
                    ON c.subject_id = s.id
                   AND c.standard_id = :std
                   AND c.sub_institute_id = :t
                 WHERE s.sub_institute_id = :t
                 GROUP BY s.id, s.subject_name
                 ORDER BY s.subject_name
                """
            ),
            {"t": tenant, "std": standard_id},
        ).fetchall()
    return {str(r[1]).strip(): int(r[0]) for r in rows}


def all_subjects(tenant: int) -> dict[str, int]:
    with SessionLocal() as db:
        rows = db.execute(
            text(
                "SELECT id, subject_name FROM subject "
                "WHERE sub_institute_id = :t ORDER BY subject_name"
            ),
            {"t": tenant},
        ).fetchall()
    return {str(r[1]).strip(): int(r[0]) for r in rows}


def chapters_for(tenant: int, standard_id: int, subject_id: int) -> list[tuple[int, str]]:
    """(chapter_number, title) in syllabus order."""
    with SessionLocal() as db:
        rows = db.execute(
            text(
                """
                SELECT sort_order, chapter_name
                  FROM chapter_master
                 WHERE sub_institute_id = :t
                   AND standard_id = :std
                   AND subject_id = :sub
                 GROUP BY sort_order, chapter_name
                 ORDER BY sort_order
                """
            ),
            {"t": tenant, "std": standard_id, "sub": subject_id},
        ).fetchall()
    out: list[tuple[int, str]] = []
    for sort_order, name in rows:
        if sort_order is None:
            continue
        out.append((int(sort_order), str(name or "").strip()))
    return out


# --------------------------------------------------------------------------
# the workbook
# --------------------------------------------------------------------------

def create_workbook(path: Path) -> None:
    """A queue workbook with nothing in it but a header."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = qs.QUEUE_SHEET_NAME

    fill = PatternFill("solid", fgColor=HEADER_FILL)
    font = Font(color="FFFFFFFF", bold=True, size=10)
    for index, column in enumerate(qs.COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=column)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[cell.column_letter].width = COLUMN_WIDTHS.get(column, 14)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(qs.COLUMNS)).column_letter}1"
    _add_validation(sheet)
    workbook.save(path)
    logger.info("Created %s", path)


def _add_validation(sheet) -> None:
    """Dropdowns on the columns where a typo would cost a night."""
    from openpyxl.worksheet.datavalidation import DataValidation

    choices = {
        "enabled": '"yes,no"',
        "board": '"CBSE,CAMBRIDGE"',
        "document_type": '"Chapter,Curriculum,Syllabus,question_bank"',
    }
    for column, formula in choices.items():
        index = qs.COLUMNS.index(column) + 1
        letter = sheet.cell(row=1, column=index).column_letter
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}5000")


def append_rows(path: Path, rows: list[qs.QueueRow]) -> int:
    """Add rows whose key is not already in the sheet. Returns how many landed."""
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    existing = {row.key for row in qs.load(path)} if path.exists() else set()
    fresh = [row for row in rows if row.key not in existing]
    if not fresh:
        return 0

    workbook = load_workbook(path)
    try:
        sheet = workbook[qs.QUEUE_SHEET_NAME]
        band = PatternFill("solid", fgColor=BAND_FILL)
        cursor = sheet.max_row + 1
        for row in fresh:
            shade = (row.chapter_number or 0) % 2 == 0
            for index, column in enumerate(qs.COLUMNS, start=1):
                value = getattr(row, column, None)
                cell = sheet.cell(row=cursor, column=index, value=value)
                if shade:
                    cell.fill = band
            cursor += 1
        workbook.save(path)
    finally:
        workbook.close()
    return len(fresh)


# NCERT publishes one PDF per chapter at a fixed path, named by the book's
# five-character code and a two-digit chapter number:
#
#     https://ncert.nic.in/textbook/pdf/jesc101.pdf   Class 10 Science ch.1
#
# so --ncert-code jesc1 fills a whole subject's pdf_url column. The code is
# printed on the book's page at ncert.nic.in; check the first link before
# running a night against it, because a wrong code yields 13 valid PDFs of the
# wrong book rather than an error.
NCERT_PDF = "https://ncert.nic.in/textbook/pdf/{code}{chapter:02d}.pdf"


def build_rows(
    *,
    board: str,
    tenant: int,
    standard_name: str,
    subject_name: str,
    chapters: list[tuple[int, str]],
    syear: int,
    document_type: str,
    ncert_code: str = "",
) -> list[qs.QueueRow]:
    return [
        qs.QueueRow(
            enabled="yes",
            board=board,
            standard=int(standard_name),
            subject_name=subject_name,
            chapter_number=number,
            document_title=title,
            document_type=document_type,
            syear=syear,
            sub_institute_id=tenant,
            pdf_url=(
                NCERT_PDF.format(code=ncert_code.strip(), chapter=number)
                if ncert_code
                else ""
            ),
            status="pending",
            attempts=0,
        )
        for number, title in chapters
    ]


# --------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--sheet", type=Path, default=DEFAULT_SHEET)
    parser.add_argument("--board", default="CBSE")
    parser.add_argument("--tenant", type=int, help="sub_institute_id; defaults to the board's")
    parser.add_argument("--standard", help="class number, e.g. 10")
    parser.add_argument("--subject", help="subject name, exactly as in the subject table")
    parser.add_argument("--all-subjects", action="store_true", help="every subject of the class")
    parser.add_argument("--syear", type=int, default=2026)
    parser.add_argument("--document-type", default="Chapter")
    parser.add_argument(
        "--blank",
        type=int,
        metavar="N",
        help="append N empty chapter rows instead of reading chapter_master",
    )
    parser.add_argument(
        "--ncert-code",
        default="",
        help="NCERT book code (e.g. jesc1) to fill pdf_url from; one subject at a time",
    )
    parser.add_argument("--list", action="store_true", help="show what the database has, then exit")
    args = parser.parse_args()

    if args.ncert_code and args.all_subjects:
        parser.error("--ncert-code names one book, so it cannot be used with --all-subjects")

    if not init_mariadb() or SessionLocal is None:
        logger.error("MariaDB is unreachable; check the MARIADB_* settings in .env")
        return 2

    tenant = _tenant(args.board, args.tenant)
    known = standards(tenant)

    if args.list:
        if not args.standard:
            logger.info("Classes for %s (tenant %s): %s", args.board, tenant, ", ".join(known))
            logger.info("Add --standard <class> to list its subjects.")
            return 0
        standard_id = known.get(str(args.standard))
        if standard_id is None:
            logger.error("Class %s not found for tenant %s. Have: %s", args.standard, tenant, ", ".join(known))
            return 2
        with_chapters = subjects_for(tenant, standard_id)
        logger.info("Class %s subjects WITH chapters in chapter_master:", args.standard)
        for name, subject_id in with_chapters.items():
            count = len(chapters_for(tenant, standard_id, subject_id))
            logger.info("   %-28s %3d chapters", name, count)
        if not with_chapters:
            logger.info("   (none -- use --blank N to make rows by hand)")
            logger.info("Subjects that exist for the tenant: %s", ", ".join(all_subjects(tenant)))
        return 0

    if not args.standard:
        parser.error("--standard is required (or use --list)")
    standard_id = known.get(str(args.standard))
    if standard_id is None:
        logger.error("Class %s not found for tenant %s. Have: %s", args.standard, tenant, ", ".join(known))
        return 2

    if not args.sheet.exists():
        create_workbook(args.sheet)

    # Which subjects to build for.
    if args.all_subjects:
        targets = subjects_for(tenant, standard_id)
        if not targets:
            logger.error("Class %s has no chapters in chapter_master; name a subject and use --blank", args.standard)
            return 2
    elif args.subject:
        pool = all_subjects(tenant)
        match = next((n for n in pool if n.lower() == args.subject.strip().lower()), None)
        if match is None:
            logger.error("Subject %r not found for tenant %s. Have: %s", args.subject, tenant, ", ".join(pool))
            return 2
        targets = {match: pool[match]}
    else:
        parser.error("pass --subject NAME or --all-subjects")

    total = 0
    for subject_name, subject_id in targets.items():
        if args.blank:
            chapters = [(n, "") for n in range(1, args.blank + 1)]
        else:
            chapters = chapters_for(tenant, standard_id, subject_id)
            if not chapters:
                logger.warning(
                    "%s has no chapter_master rows for class %s; skipping "
                    "(use --blank N to add empty rows)",
                    subject_name,
                    args.standard,
                )
                continue

        rows = build_rows(
            board=args.board,
            tenant=tenant,
            standard_name=str(args.standard),
            subject_name=subject_name,
            chapters=chapters,
            syear=args.syear,
            document_type=args.document_type,
            ncert_code=args.ncert_code,
        )
        added = append_rows(args.sheet, rows)
        if args.ncert_code and rows:
            logger.info("   pdf_url filled from %s -- check %s opens", args.ncert_code, rows[0].pdf_url)
        total += added
        logger.info(
            "%-28s %2d chapter(s) in the database, %2d new row(s) added",
            subject_name,
            len(chapters),
            added,
        )

    logger.info("")
    logger.info("Sheet: %s", args.sheet)
    logger.info("%d row(s) added. Fill in the pdf_url column, then run:", total)
    logger.info("   python -m scripts.run_extraction_queue --list")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
